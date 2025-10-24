import json
import traceback
from docx import Document
from typing import Dict, Any
from pathlib import Path


def get_nested_value(data: Dict[str, Any], key_path: str) -> str:
    keys = key_path.split('.')
    value = data
    
    try:
        for key in keys:
            if isinstance(value, dict):
                value = value.get(key, "")
            else:
                return ""
        
        if isinstance(value, dict):
            return json.dumps(value, indent=2, ensure_ascii=False)
        
        return str(value) if value is not None else ""
    
    except (KeyError, TypeError, AttributeError):
        return ""


def generate_pep_document(template_path: str, data: Dict[str, Any], output_path: str) -> bool:
    # Extraction directe du champ 'body' du format n8n
    data = data[0]["body"]
    placeholder_mapping = {
        "{{GENERALITES}}": "descriptif_projet.generalites",
        "{{JUSTIFICATION}}": "descriptif_projet.justification",
        "{{ASPECT_CONTRACTUEL}}": "descriptif_projet.aspect_contractuel",
        "{{SCOPE_PROJET}}": "descriptif_projet.description_detaillee.scope_projet",
        "{{BASE_DESIGN}}": "descriptif_projet.description_detaillee.base_design",
        "{{CONTRAINTES_PRINCIPALES}}": "descriptif_projet.description_detaillee.contraintes_principales",
        "{{DESCRIPTION_DETAILLEE_INSTALLATIONS_ET_SOLUTIONS_RETENUES}}": "descriptif_projet.description_detaillee.description_detaillee_installations_et_solutions_retenues",
        "{{POINTS_EN_ATTENTE}}": "descriptif_projet.description_detaillee.points_en_attente",
        "{{ORGANISATION_EQUIPE_PROJET}}": "organisation_equipe_projet",
        "{{DOCUMENTS_CLIENT_DE_REFERENCE}}": "documents_reference.documents_client_de_reference",
        "{{DOCUMENTS_CONTRACTUELS}}": "documents_reference.documents_contractuels",
        "{{DOCUMENTS_REFERENCE}}": "documents_reference.documents_reference",
        "{{DOCUMENTS_INTERNES}}": "documents_reference.documents_internes",
        "{{ORGANISATION_GENERALE}}": "gestion_projet.organisation_generale",
        "{{MATRICE_RESPONSABILITES}}": "gestion_projet.matrice_responsabilites",
        "{{JALONS_PRINCIPAUX}}": "gestion_projet.jalons_principaux",
        "{{REUNIONS}}": "gestion_projet.reunions",
        "{{AVANCEMENT_PHYSIQUE}}": "gestion_projet.controle_projet.avancement_physique",
        "{{PLANNING}}": "gestion_projet.controle_projet.planning",
        "{{CONTROLE_COUTS}}": "gestion_projet.controle_projet.controle_couts",
        "{{RISK_MANAGEMENT}}": "gestion_projet.controle_projet.risk_management",
        "{{REPORTING}}": "gestion_projet.reporting",
        "{{GESTION_SCOPE_ET_MODIFICATIONS}}": "gestion_projet.gestion_scope_et_modifications",
        "{{COMMUNICATIONS}}": "gestion_projet.communications",
        "{{GESTION_DE_DOCUMENTATION}}": "gestion_projet.gestion_de_documentation",
        "{{CERTIFICATION_BUREAU_DE_CONTROLE}}": "certification_bureau_de_controle",
        "{{ACTIVITES_ET_DONNEES_STRATEGIQUES}}": "ingenierie.activites_et_donnees_strategiques",
        "{{LISTE_DE_LIVRABLES _ET_ACTIVITES_REPARTITION_DES_ROLES_ET_RESPONSABILITES}}": "ingenierie.liste_de_livrables_et_activites_repartition_des_roles_et_responsabilites",
        "{{EXCLUSIONS}}": "ingenierie.exclusions",
        "{{INTERFACES_ET_LIMITES_DES_PRESTATIONS}}": "ingenierie.interfaces_et_limites_des_prestations",
        "{{BATTERRIE _LIMITES_TECHNIQUES}}": "ingenierie.batterie_limites_techniques",
        "{{INTERFACE _DANS_ROLES_ET_REPONSABILITES}}": "ingenierie.interface_dans_roles_et_responsabilites",
        "{{PLANNING_DES_ETUDES}}": "ingenierie.planning_des_etudes",
        "{{MAQUETTE_3D_CAD_OUTILS_ET_LOGICIELS}}": "ingenierie.maquette_3d_cad_outils_et_logiciels",
        "{{RISQUES_INGENIERIE_IDENTIFIES_ET_PLAN_DE_MIGRATION}}": "ingenierie.risques_ingenierie_identifies_et_plan_de_mitigation",
        "{{PLAN_DE_VERIFICATION_DES_ETUDES}}": "ingenierie.plan_de_verification_des_etudes",
        "{{PLAN_DE_REVUES_INGENIERIE}}": "ingenierie.plan_de_revues_ingenierie",
        "{{APPROVISONNEMENTS_PROCUREMENT}}": "approvisionnements_procurement.description_generale",
        "{{TUYAUTERIES}}": "approvisionnements_procurement.tuyauteries",
        "{{INSTRUMENTATIONS}}": "approvisionnements_procurement.instrumentations",
        "{{AUTOMATISATISMES_SECURITE}}": "approvisionnements_procurement.automatismes_securite",
        "{{MECANIQUES_ET_EQUIPEMENTS}}": "approvisionnements_procurement.mecaniques_et_equipements",
        "{{ELECTRICITE}}": "approvisionnements_procurement.electricite",
        "{{GENIE_CIVIL_STRUCTURE}}": "approvisionnements_procurement.genie_civil_structure",
        "{{EXPEDITING_ET_RECEPTION_MATERIEL}}": "approvisionnements_procurement.expediting_et_reception_materiel.description",
        "{{RECEPTION_DU _MATERIEL}}": "approvisionnements_procurement.expediting_et_reception_materiel.reception_du_materiel",
        "{{FACTORY_ACCEPTANCE_TEST}}": "approvisionnements_procurement.expediting_et_reception_materiel.factory_acceptance_test",
        "{{SITE_ACCEPTANCE_TEST}}": "approvisionnements_procurement.expediting_et_reception_materiel.site_acceptance_test",
        "{{TEST_DE_PERFORMANCE_GARANTIES}}": "approvisionnements_procurement.expediting_et_reception_materiel.test_de_performance_garanties",
        "{{MARCHES_DE_TRAVAUX}}": "marches_de_travaux",
        "{{CONSTRUCTION_GENERALITES}}": "construction.generalites",
        "{{INSTALLATIONS_TEMPORAIRES}}": "construction.installations_temporaires",
        "{{HSE_CHANTIER}}": "construction.hse_chantier",
        "{{PIPING_CALO_ECHAFAUDAGES}}": "construction.piping_calo_echafaudages",
        "{{EIA}}": "construction.eia",
        "{{PROCESS_CONTROL_AUTOMATISME}}": "construction.process_control_automatisme",
        "{{AUTRES_LEVERAGE_LOURD_MONTAGE_MECANIQUE}}": "construction.autres_levage_lourd_montage_mecanique",
        "{{MISES_A_DISPOSITIONS}}": "construction.mises_a_dispositions",
        "{{QUALITES_DES_TRAVAUX}}": "construction.qualites_des_travaux",
        "{{CONSIGNATIONS_ET_DECONSIGNATIONS}}": "construction.consignations_et_deconsignations",
        "{{MECHANICAL_COMPLETION}}": "construction.mechanical_completion",
        "{{PRECOM_COMMISSIONING_MISE EN SERVICE}}": "precom_commissioning_mise_en_service",
        "{{PIECES_DE_RECHANGES}}": "pieces_de_rechanges",
        "{{DESAFFECTION_DU_MATERIEL}}": "desaffection_du_materiel",
        "{{FORMATIONS_DU_PERSONNEL_CLIENT}}": "formations_du_personnel_client",
        "{{AUTORISATION_D'EXPLOITER _PERMIS_DE_CONSTRUIRE}}": "autorisation_exploiter_permis_de_construire",
    }
    
    try:
        doc = Document(template_path)
        
        # Remplacer dans les paragraphes
        for paragraph in doc.paragraphs:
            for placeholder, json_path in placeholder_mapping.items():
                if placeholder in paragraph.text:
                    value = get_nested_value(data, json_path)
                    paragraph.text = paragraph.text.replace(placeholder, value)
        
        # Remplacer dans les tableaux
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for placeholder, json_path in placeholder_mapping.items():
                            if placeholder in paragraph.text:
                                value = get_nested_value(data, json_path)
                                paragraph.text = paragraph.text.replace(placeholder, value)
        
        doc.save(output_path)
        
        return True
        
    except Exception as e:
        traceback.print_exc()
        return False